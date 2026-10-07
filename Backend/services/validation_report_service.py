"""Automated Snowflake-vs-BigQuery Power BI validation DOCX.

Produces the same document the team fills in by hand (see
"PBI Report Validation - Template"), straight from the persisted run
artifacts:

    1. Dashboard & Environment Details
    2. Semantic Model Refresh Check          (+ 2.1 GCP / 2.2 Snowflake shots)
    3. Filter Validation Log                 (landing page, default load)
    4. Visual & Functional Validation Log    (same visual on both dashboards?)
    5. Data Export Validation Log            (can the table/matrix export?)
    6. Summary & Sign-off

Rules this module follows
-------------------------
* Nothing is invented. A value the capture did not record is left BLANK
  (never "Not Available", "None" or "N/A"), exactly like an untouched
  template cell.
* Data sources:  DOM -> filters, refresh date, navigation, export results,
  visual inventory.   Gemini -> dashboard name / refresh date fallback.
  Playwright -> screenshots.   Tester -> persisted tester name when available.
* Status cells are real Word dropdown content controls (Pass / Fail /
  In-Progress ...), so a human can still change them in Word.
* Mismatch detail is NOT printed here - the front end already shows it. This
  document only records what matched / what was checked.

Source = Snowflake (SpartanNash), target = GCP / BigQuery (Tredence). Change
``_SOURCE_LABEL`` / ``_TARGET_LABEL`` if that mapping is ever reversed.

Entry point (unchanged):

    path = await build_validation_report(run_id)
"""

from __future__ import annotations

import asyncio
import logging
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from services.dashboard_inventory_service import _classify_dom_visual, _kpi_key
from services.run_artifact_store import load_run_status, load_validation_capture
from utils.config import REPORT_DIR, SCREENSHOT_DIR


logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

_SOURCE_LABEL = "Snowflake"
_TARGET_LABEL = "GCP"
_TARGET_REFRESH_HEADER = "BigQuery"  # the template's refresh-table column title

_TESTER_NAME = ""  # fallback only; never invent a tester name
_DOC_SUBTITLE = "Snowflake vs. BigQuery Power BI Report Validation"
_DEFAULT_TITLE = "Dashboard Validation Report"

# Optional footer logos (the template has C&S on the left, SpartanNash on the
# right). Leave as None to omit; a missing file is silently skipped.
_LOGO_LEFT: Path | None = None
_LOGO_RIGHT: Path | None = None

_LANDING_ACTION = "Landing Page (default load, no filters changed)"
_LANDING_EXPECTED = (
    "Every tab/visual that reads this filter recalculates identically on "
    "both sources"
)
_VISUAL_ACTION = "Landing Page (default load, no filters changed)"
_VISUAL_EXPECTED = "Visual is present on both dashboards"

_FONT = "Arial"
_ACCENT = RGBColor(0x1F, 0x38, 0x64)
_GRAY = RGBColor(0x59, 0x59, 0x59)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_HEADER_FILL = "1F3864"
_LABEL_FILL = "F2F2F2"
_BORDER = "BFBFBF"

_PORTRAIT_MARGIN_IN = 1.0
_IMAGE_MAX_W_IN = 6.5
_IMAGE_MAX_H_IN = 4.3

_BODY_PT = 10.0
_TABLE_PT = 9.0

_PASS = "PASS"
_FAIL = "FAIL"
_UNCERTAIN = "UNCERTAIN"
_NOT_AVAILABLE_LABEL = ""

_PERSISTED_REFRESH_KEYWORDS = (
    "refresh",
    "last_updated",
    "updated_at",
    "updated_as_of",
    "data_as_of",
)
_CHECKS = ("kpis", "visuals", "filters", "buttons")

_IN_PROGRESS_RUN_STATES = {
    "running", "in progress", "queued", "pending", "processing",
    "started", "capturing", "comparing", "initializing",
}
_FAILED_RUN_STATES = {"failed", "failure", "error"}


class ValidationReportNotFound(Exception):
    """Raised when no persisted capture exists for a run_id."""


class ValidationReportError(Exception):
    """Raised when the DOCX report cannot be generated."""


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------


def _clean(value) -> str:
    """Stripped text, or "" - blank is the only 'unknown' this report uses."""
    if value is None or isinstance(value, (list, tuple, dict)):
        return ""
    return str(value).strip()


def _safe_filename_dashboard(name: str) -> str:
    text = _clean(name) or "Dashboard"
    text = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", text)
    text = re.sub(r"\s+", "_", text)
    text = re.sub(r"_+", "_", text)
    return text.strip(" ._") or "Dashboard"


def report_filename(run_id: str, dashboard_name: str) -> str:
    """``<Dashboard_Name>_<run_id>.docx``."""
    return f"{_safe_filename_dashboard(dashboard_name)}_{run_id}.docx"


def _resolve_screenshot_path(recorded: str | None) -> Path | None:
    """Existing screenshot file for a recorded path (handles stale roots)."""
    if not recorded:
        return None
    candidate = Path(recorded)
    if candidate.exists():
        return candidate
    rerooted = Path(
        str(recorded)
        .replace("\\output\\", "\\Backend\\output\\")
        .replace("/output/", "/Backend/output/")
    )
    if rerooted.exists():
        return rerooted
    matches = sorted(Path(SCREENSHOT_DIR).rglob(candidate.name))
    return matches[0] if matches else None


def _date_only(value) -> str:
    """dd-mm-yyyy from an ISO-ish timestamp; seconds/time are dropped."""
    text = _clean(value)
    if not text:
        return ""
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).strftime("%d-%m-%Y")
    except ValueError:
        pass
    match = re.search(r"(\d{4})-(\d{2})-(\d{2})", text)
    return f"{match.group(3)}-{match.group(2)}-{match.group(1)}" if match else ""


# ---------------------------------------------------------------------------
# Status presentation (mirrors the comparison engine's vocabulary)
# ---------------------------------------------------------------------------

_PASS_WORDS = {
    "match", "matched", "table matched", "pass", "passed", "success",
    "succeeded", "verified", "stable", "completed", "applied", "confirmed",
}
_FAIL_WORDS = {
    "mismatch", "mismatched", "table mismatch", "table mismatched",
    "missing in source", "missing in target", "fail", "failed", "failure",
    "error", "unstable", "render not confirmed",
}
_UNCERTAIN_WORDS = {
    "needs review", "uncertain", "mixed", "partial", "inconclusive", "ambiguous",
}
_UNAVAILABLE_WORDS = {
    "not compared", "table not compared", "unavailable", "not available",
    "no data", "not tested", "skipped",
}


def _status_label(status, reason=None) -> str:
    """Normalize a persisted status for the DOCX.

    Missing/unknown information stays blank.  We never turn an unknown value
    into PASS, FAIL, or the literal text "Not Available".
    """
    text = _clean(status)
    if not text:
        return ""
    lowered = " ".join(text.casefold().replace("_", " ").split())
    if lowered in _PASS_WORDS:
        return _PASS
    if lowered in _FAIL_WORDS:
        return _FAIL
    if lowered in _UNCERTAIN_WORDS:
        return _UNCERTAIN
    if lowered in _UNAVAILABLE_WORDS or lowered in {"not available", "n/a", "none"}:
        return ""
    return _UNCERTAIN


def _iter_comparison_items(response: dict):
    comparison = response.get("comparison") or {}
    page_comparisons = comparison.get("page_comparisons") or []
    if page_comparisons:
        for page in page_comparisons:
            for kind in _CHECKS:
                for item in page.get(kind) or []:
                    if isinstance(item, dict):
                        yield kind, item
            for item in (page.get("tables") or {}).get("comparisons") or []:
                if isinstance(item, dict):
                    yield "tables", item
        return
    for kind in _CHECKS:
        for item in comparison.get(kind) or []:
            if isinstance(item, dict):
                yield kind, item
    tables = comparison.get("tables") or {}
    if isinstance(tables, dict):
        for item in tables.get("comparisons") or []:
            if isinstance(item, dict):
                yield "tables", item


def _overall_status(response: dict, run_status: dict) -> str:
    """Return the report-level dropdown value without inventing a result."""
    lifecycle = " ".join(
        _clean((run_status or {}).get("status")).casefold().replace("_", " ").split()
    )
    if lifecycle in _IN_PROGRESS_RUN_STATES:
        return "In-Progress"
    if lifecycle in _FAILED_RUN_STATES:
        return "Fail"

    # Prefer an explicit persisted overall result when one exists.
    candidates = [
        (response or {}).get("overall_status"),
        (response or {}).get("overall_result"),
        ((response or {}).get("summary") or {}).get("overall_result"),
        ((response or {}).get("comparison") or {}).get("status"),
    ]
    for candidate in candidates:
        label = _status_label(candidate)
        if label == _FAIL:
            return "Fail"
        if label == _PASS:
            return "Pass"
        if label == _UNCERTAIN:
            return "In-Progress"

    labels = [
        _status_label(item.get("status"), item.get("reason"))
        for _kind, item in _iter_comparison_items(response)
    ]
    labels = [label for label in labels if label]
    if _FAIL in labels:
        return "Fail"
    if labels and all(label == _PASS for label in labels):
        return "Pass"
    if labels:
        return "In-Progress"
    return ""


# ---------------------------------------------------------------------------
# Page pairs
# ---------------------------------------------------------------------------


@dataclass
class PagePair:
    page_name: str
    source: dict = field(default_factory=dict)  # empty => page missing on source
    target: dict = field(default_factory=dict)  # empty => page missing on target
    comparison: dict | None = None
    source_image: Path | None = None
    target_image: Path | None = None


def _split_executions(document: dict) -> tuple[list, list]:
    groups = document.get("executions_by_dashboard") or []
    source = groups[0] if len(groups) > 0 else []
    target = groups[1] if len(groups) > 1 else []
    return source or [], target or []


def _execution_page_name(execution: dict) -> str:
    dashboard = execution.get("dashboard") or {}
    metadata = (execution.get("visual_data") or {}).get("metadata") or {}
    return _clean(dashboard.get("page_name") or metadata.get("page_name"))


def _shot(execution: dict | None) -> Path | None:
    if not execution:
        return None
    return _resolve_screenshot_path(
        (execution.get("metrics") or {}).get("screenshot_path")
    )


def _collect_pages(document: dict, response: dict) -> list[PagePair]:
    """Every page seen on either dashboard, in source order then target-only.

    A page that exists on only one side is kept (with an empty execution on
    the other side) so the filter log can record it as a failed landing check
    instead of silently dropping it.
    """
    source_executions, target_executions = _split_executions(document)
    comparisons = {
        item.get("page_name"): item
        for item in (response.get("comparison") or {}).get("page_comparisons", [])
    }

    target_by_name: dict[str, dict] = {}
    for execution in target_executions:
        target_by_name.setdefault(_execution_page_name(execution), execution)

    pages: list[PagePair] = []
    used: set[str] = set()
    for index, source in enumerate(source_executions):
        name = _execution_page_name(source)
        target = target_by_name.get(name)
        if target is None and not name and index < len(target_executions):
            target = target_executions[index]  # unnamed single-page report
        if target is not None:
            used.add(_execution_page_name(target))
        pages.append(_make_pair(name, source, target or {}, comparisons))

    for target in target_executions:
        name = _execution_page_name(target)
        if name in used:
            continue
        pages.append(_make_pair(name, {}, target, comparisons))
    return pages


def _make_pair(name: str, source: dict, target: dict, comparisons: dict) -> PagePair:
    comparison = comparisons.get(name)
    if comparison is None and len(comparisons) == 1 and not name:
        comparison = next(iter(comparisons.values()))
    return PagePair(
        page_name=name or "Landing Page",
        source=source,
        target=target,
        comparison=comparison,
        source_image=_shot(source),
        target_image=_shot(target),
    )


# ---------------------------------------------------------------------------
# Section 1 / 2 data: dashboard name + refresh stamps
# ---------------------------------------------------------------------------


def _dom_report_name(metrics: list) -> str:
    """Report title from the browser tab titles ('<Page> - <Report> - Power BI')."""
    candidates: dict[str, int] = {}
    for metric in metrics or []:
        if not isinstance(metric, dict):
            continue
        parts = _clean(metric.get("page_title")).split(" - ")
        if parts[-1].strip().casefold() not in ("power bi", "microsoft power bi"):
            continue
        if len(parts) >= 3:
            segment = parts[-2].strip()
        elif len(parts) == 2:
            segment = parts[0].strip()
        else:
            continue
        if segment:
            candidates[segment] = candidates.get(segment, 0) + 1
    return max(candidates, key=candidates.get) if candidates else ""


def _report_display_name(source_name: str, target_name: str, metrics: list) -> str:
    """Resolve the actual dashboard name; never synthesize an A-vs-B name."""
    return _dom_report_name(metrics) or _clean(source_name) or _clean(target_name)


def resolve_dashboard_name(source_name: str, target_name: str, metrics: list) -> str:
    """Public resolver used by the API response (behaviour unchanged)."""
    return _report_display_name(source_name, target_name, metrics)


def _looks_like_stamp(value: str) -> bool:
    # A refresh timestamp always carries digits; this rejects things like
    # {"refresh_status": "success"} being mistaken for a date.
    return bool(re.search(r"\d", value))


_LABEL_PREFIX = re.compile(
    r"^\s*(?:data\s+)?(?:last\s+)?(?:refresh(?:ed)?|updated)(?:\s+(?:on|at))?\s*[:\-]?\s*",
    re.IGNORECASE,
)


def _tidy_stamp(value: str) -> str:
    stripped = _LABEL_PREFIX.sub("", value).strip()
    return stripped or value.strip()


def _scan_refresh(obj) -> str:
    if isinstance(obj, dict):
        for key, value in obj.items():
            lowered = str(key).casefold()
            if (
                any(token in lowered for token in _PERSISTED_REFRESH_KEYWORDS)
                and isinstance(value, str)
                and value.strip()
                and _looks_like_stamp(value)
            ):
                return value.strip()
        for value in obj.values():
            found = _scan_refresh(value)
            if found:
                return found
    elif isinstance(obj, list):
        for item in obj:
            found = _scan_refresh(item)
            if found:
                return found
    return ""


def _dom_refresh(executions: list) -> str:
    for execution in executions or []:
        for container in (
            execution,
            execution.get("dashboard") or {},
            execution.get("metrics") or {},
            execution.get("extraction") or {},
            execution.get("visual_data") or {},
        ):
            found = _scan_refresh(container)
            if found:
                return found
    return ""


_SOURCE_TOKENS = ("source", "spartnash", "snowflake")
_TARGET_TOKENS = ("target", "trendence", "tredence", "gcp", "bigquery")


def _side_of(key: str) -> str | None:
    if any(token in key for token in _SOURCE_TOKENS):
        return "source"
    if any(token in key for token in _TARGET_TOKENS):
        return "target"
    return None


def _scan_ai_refresh(obj, want: str, context: str | None = None) -> str:
    """Gemini fallback: refresh-ish key that belongs to ``want`` (source/target).

    A key belongs to a side when the key itself or an ancestor key names it
    (``source_last_refresh``, ``spartnash.last_refresh`` ...).
    """
    if isinstance(obj, dict):
        for key, value in obj.items():
            lowered = str(key).casefold()
            side = _side_of(lowered) or context
            if (
                side == want
                and isinstance(value, str)
                and _looks_like_stamp(value)
                and any(token in lowered for token in _PERSISTED_REFRESH_KEYWORDS)
            ):
                return value.strip()
        for key, value in obj.items():
            found = _scan_ai_refresh(value, want, _side_of(str(key).casefold()) or context)
            if found:
                return found
    elif isinstance(obj, list):
        for item in obj:
            found = _scan_ai_refresh(item, want, context)
            if found:
                return found
    return ""


def _ai_refresh_stamp(ai: dict, want: str) -> str:
    validation = (ai or {}).get("refresh_validation")
    if not isinstance(validation, (list, dict)):
        validation = ((ai or {}).get("ai_analysis") or {}).get("refresh_validation")

    if isinstance(validation, dict):
        side = validation.get(want) or {}
        if isinstance(side, dict):
            value = side.get("refresh_datetime") or side.get("refresh_date")
            if isinstance(value, str) and value.strip() and _looks_like_stamp(value):
                return value.strip()

    if isinstance(validation, list):
        value_key = "source_value" if want == "source" else "target_value"
        for item in validation:
            if not isinstance(item, dict):
                continue
            value = item.get(value_key)
            if isinstance(value, str) and value.strip() and _looks_like_stamp(value):
                return value.strip()
    return ""


def _refresh_screenshot(executions: list, side: str) -> Path | None:
    """Only return a screenshot explicitly recorded for refresh validation."""
    side_tokens = _SOURCE_TOKENS if side == "source" else _TARGET_TOKENS

    def walk(obj):
        if isinstance(obj, dict):
            for key, value in obj.items():
                lowered = str(key).casefold()
                if (
                    "screenshot" in lowered
                    and "refresh" in lowered
                    and isinstance(value, str)
                    and value
                ):
                    shot = _resolve_screenshot_path(value)
                    if shot:
                        return shot
                found = walk(value)
                if found:
                    return found
        elif isinstance(obj, list):
            for item in obj:
                found = walk(item)
                if found:
                    return found
        return None

    # Prefer an explicitly side-labelled refresh screenshot.
    for execution in executions or []:
        if any(token in str(execution).casefold() for token in side_tokens):
            shot = walk(execution)
            if shot:
                return shot
    for execution in executions or []:
        shot = walk(execution)
        if shot:
            return shot
    return None


def _refresh_stamp(executions: list, ai: dict, want: str) -> str:
    found = (
        _dom_refresh(executions)
        or _scan_ai_refresh(ai or {}, want)
        or _ai_refresh_stamp(ai or {}, want)
    )
    return _tidy_stamp(found) if found else ""


# ---------------------------------------------------------------------------
# Section 3: filter (landing page) rows
# ---------------------------------------------------------------------------


def _page_filter_result(page: PagePair) -> str | None:
    """'pass' / 'fail' / None (blank) for a landing-page default load."""
    if not page.source or not page.target:
        return "fail"  # the page exists on only one dashboard
    if page.comparison is None:
        return None
    labels = [
        _status_label(item.get("status"), item.get("reason"))
        for item in page.comparison.get("filters") or []
        if isinstance(item, dict)
    ]
    if _FAIL in labels:
        return "fail"
    if any(label in (_UNCERTAIN, _NOT_AVAILABLE_LABEL) for label in labels):
        return None
    return "pass"


def _iter_slicer_scenarios(document: dict, response: dict) -> list[dict]:
    """Find persisted slicer/filter execution evidence without assuming one schema."""
    candidates = []
    for container in (
        document,
        response,
        (document or {}).get("validation") or {},
        (document or {}).get("results") or {},
        (response or {}).get("comparison") or {},
        (response or {}).get("ai_analysis") or {},
    ):
        for key in ("slicer_scenarios", "filter_scenarios", "slicer_results"):
            value = container.get(key) if isinstance(container, dict) else None
            if isinstance(value, list):
                candidates.extend(x for x in value if isinstance(x, dict))
    return candidates


def _scenario_result(scenario: dict) -> str:
    """Use runtime application/render evidence first, then comparison result."""
    explicit = _clean(scenario.get("result") or scenario.get("status"))
    if explicit:
        label = _status_label(explicit, scenario.get("details") or scenario.get("reason"))
        if label in {_PASS, _FAIL}:
            # Runtime evidence below can still override a false PASS.
            if label == _PASS:
                runtime_fields = (
                    "source_applied", "target_applied",
                    "source_render_stable", "target_render_stable",
                )
                present = [f for f in runtime_fields if f in scenario]
                if present and not all(bool(scenario.get(f)) for f in present):
                    return _FAIL
            return label

    applied_fields = [f for f in ("source_applied", "target_applied") if f in scenario]
    render_fields = [
        f for f in ("source_render_stable", "target_render_stable")
        if f in scenario
    ]

    if applied_fields and not all(bool(scenario.get(f)) for f in applied_fields):
        return _FAIL
    if render_fields and not all(bool(scenario.get(f)) for f in render_fields):
        return _FAIL
    if applied_fields or render_fields:
        return _PASS

    return ""


def _filter_rows(pages: list[PagePair], document: dict, response: dict) -> list[list]:
    """Return exactly: Page Name | Action / Filter Applied | Result.

    Landing/default-load is always the first row when at least one page exists.
    Persisted slicer scenarios are then appended.  No internal render timings
    are printed.
    """
    rows: list[list] = []

    for page in pages:
        result = _page_filter_result(page)
        rows.append([
            page.page_name,
            "Landing Page - No Filter Applied",
            _choose(OPT_PASS_FAIL, "Result", result),
        ])

    scenarios = _iter_slicer_scenarios(document, response)
    for scenario in scenarios:
        page_name = _clean(
            scenario.get("page_name")
            or scenario.get("page")
            or scenario.get("tab_name")
            or "Landing Page"
        )
        slicer = _clean(
            scenario.get("slicer")
            or scenario.get("slicer_name")
            or scenario.get("filter")
            or scenario.get("filter_name")
            or scenario.get("item_name")
        )
        value = _clean(
            scenario.get("value")
            or scenario.get("selected_value")
            or scenario.get("filter_value")
        )
        action = slicer
        if value:
            action = f"{slicer} = {value}" if slicer else value
        if not action:
            action = "Filter Applied"
        rows.append([
            page_name,
            action,
            _choose(OPT_PASS_FAIL, "Result", _scenario_result(scenario)),
        ])

    # If the runtime did not persist slicer scenarios, the AI comparison JSON
    # can still describe which filters were checked.  Use it only as a
    # comparison fallback; never pretend it proves that a browser click
    # happened or that the dashboard rendered.
    if not scenarios:
        ai_filters = (response or {}).get("filter_validation")
        if not isinstance(ai_filters, list):
            ai_filters = ((response or {}).get("ai_analysis") or {}).get(
                "filter_validation"
            )
        if isinstance(ai_filters, list):
            for item in ai_filters:
                if not isinstance(item, dict):
                    continue
                page_name = _clean(
                    item.get("page_name") or item.get("page") or "Landing Page"
                )
                name = _clean(item.get("item_name") or item.get("filter_name"))
                source_value = _clean(item.get("source_value"))
                target_value = _clean(item.get("target_value"))
                if source_value and target_value and source_value == target_value:
                    action = f"{name} = {source_value}" if name else source_value
                elif source_value or target_value:
                    action = (
                        f"{name} ({_SOURCE_LABEL}: {source_value}; "
                        f"{_TARGET_LABEL}: {target_value})"
                    )
                else:
                    action = name or "Filter Checked"
                result = _result_from_item(item)
                rows.append([
                    page_name,
                    action,
                    _choose(OPT_PASS_FAIL, "Result", result),
                ])

    return rows


# ---------------------------------------------------------------------------
# Section 4: visual presence
# ---------------------------------------------------------------------------

_KIND_LABELS = {
    "kpi": "KPI card", "table": "Table", "matrix": "Matrix",
    "bar": "Bar chart", "line": "Line chart", "area": "Area chart",
    "pie": "Pie chart", "scatter": "Scatter chart", "map": "Map",
    "card": "Card", "gauge": "Gauge", "treemap": "Treemap",
    "funnel": "Funnel", "waterfall": "Waterfall", "other": "Visual",
}
_GENERATED_TITLE = re.compile(r"(?:^|_)(?:table|matrix)_\d+$", re.IGNORECASE)


def _real_title(value) -> str:
    """A caption the report itself gave the visual (not a generated one)."""
    title = _clean(value)
    return "" if (not title or _GENERATED_TITLE.search(title)) else title


def _visual_inventory(execution: dict) -> list[dict]:
    """Charts, KPI cards, tables and matrices on one page (slicers excluded)."""
    visual_data = (execution or {}).get("visual_data") or {}
    table_visuals = visual_data.get("table_visuals") or visual_data.get("table_exports") or []

    items: list[dict] = []
    for visual in visual_data.get("visuals") or []:
        if not isinstance(visual, dict) or visual.get("is_loading_placeholder"):
            continue
        kind = _classify_dom_visual(visual)
        if kind == "slicer":
            continue
        if table_visuals and kind in ("table", "matrix"):
            continue  # table_visuals is the authoritative table list
        items.append({"kind": kind, "title": _real_title(visual.get("title"))})

    for visual in table_visuals:
        if isinstance(visual, dict):
            items.append(
                {
                    "kind": "matrix" if visual.get("is_matrix") else "table",
                    "title": _real_title(visual.get("title")),
                }
            )

    seen: set[str] = set()
    for card in visual_data.get("kpi_cards") or []:
        key = _kpi_key(card.get("name"))
        if key and key not in seen:
            seen.add(key)
            items.append({"kind": "kpi", "title": _clean(card.get("name"))})
    return items


def _norm(text: str) -> str:
    return " ".join(text.casefold().split())


def _result_from_item(item: dict | None) -> str:
    if not item:
        return ""
    result = _clean(item.get("result") or item.get("status"))
    label = _status_label(result, item.get("details") or item.get("reason"))
    return label


def _ai_visual_items(response: dict) -> list[dict]:
    """Read the AI report format when it is present.

    The AI report is evidence for visual/value comparison; the DOM inventory
    remains authoritative for visual presence/type.
    """
    items = (response or {}).get("visual_functional_validation")
    if isinstance(items, list):
        return [x for x in items if isinstance(x, dict)]
    items = ((response or {}).get("ai_analysis") or {}).get("visual_functional_validation")
    if isinstance(items, list):
        return [x for x in items if isinstance(x, dict)]
    return []


def _visual_rows(page: PagePair, response: dict | None = None) -> list[list]:
    """Return exactly: Page Name | Visual Name / Visual Type | Result."""
    source = _visual_inventory(page.source)
    target = _visual_inventory(page.target)
    ai_items = _ai_visual_items(response or {})

    # Index AI visual results by normalized item name.
    ai_by_name: dict[str, dict] = {}
    for item in ai_items:
        name = _norm(_clean(item.get("item_name")))
        if name:
            ai_by_name[name] = item

    kinds: list[str] = []
    for item in [*source, *target]:
        if item["kind"] not in kinds:
            kinds.append(item["kind"])

    rows: list[list] = []
    for kind in kinds:
        label = _KIND_LABELS.get(kind, "Visual")
        source_items = [i for i in source if i["kind"] == kind]
        target_left = list(target for target in target if target["kind"] == kind)

        source_left: list[dict] = []
        pairs: list[tuple[dict, dict]] = []
        for item in source_items:
            key = _norm(item["title"])
            hit = next(
                (t for t in target_left if key and _norm(t["title"]) == key),
                None,
            )
            if hit is None:
                source_left.append(item)
            else:
                target_left.remove(hit)
                pairs.append((item, hit))

        while source_left and target_left:
            pairs.append((source_left.pop(0), target_left.pop(0)))

        total = len(pairs) + len(source_left) + len(target_left)

        def display_name(index: int, item: dict) -> str:
            title = _real_title(item.get("title"))
            if title:
                return f"{title} ({label})"
            return f"{label} {index}" if total > 1 else label

        position = 0
        for source_item, target_item in pairs:
            position += 1
            name = display_name(position, source_item)
            ai_item = ai_by_name.get(_norm(_real_title(source_item.get("title"))))
            if ai_item is None:
                ai_item = ai_by_name.get(_norm(_real_title(target_item.get("title"))))
            result = _result_from_item(ai_item)
            if not result:
                result = _PASS
            rows.append([page.page_name, name, _choose(OPT_PASS_FAIL, "Result", result)])

        for item in source_left:
            position += 1
            name = f"{display_name(position, item)} - only on {_SOURCE_LABEL}"
            rows.append([page.page_name, name, _choose(OPT_PASS_FAIL, "Result", _FAIL)])

        for item in target_left:
            position += 1
            name = f"{display_name(position, item)} - only on {_TARGET_LABEL}"
            rows.append([page.page_name, name, _choose(OPT_PASS_FAIL, "Result", _FAIL)])

    return rows


# ---------------------------------------------------------------------------
# Section 5: table / matrix export capability
# ---------------------------------------------------------------------------


def _normalise_title(title) -> str:
    return " ".join(str(title or "").casefold().split())


def _pairing_key(record: dict) -> str:
    comparison_key = _clean((record or {}).get("comparison_key"))
    if comparison_key:
        return comparison_key.casefold()
    return _normalise_title((record or {}).get("title"))


def _is_browser_gone(error) -> bool:
    text = str(error or "").casefold()
    return any(
        marker in text
        for marker in (
            "has been closed", "targetclosederror", "target page closed",
            "browser has been closed", "connection closed",
        )
    )


def _export_state(export: dict | None, present: bool) -> str:
    """yes / no / unknown / absent for one side of one table."""
    if export is None:
        return "no" if present else "absent"
    if export.get("status") in ("downloaded", "success"):
        return "yes"
    if export.get("export_outcome") == "browser_closed" or _is_browser_gone(export.get("error")):
        return "unknown"  # the browser died - says nothing about the visual
    return "no"


_STATE_TEXT = {
    "yes": "Yes", "no": "No", "unknown": "Not confirmed", "absent": "Table not found",
}


def _export_rows(pages: list[PagePair]) -> list[list]:
    """Final export outcome only.

    Retry logic belongs to the browser/export executor.  This function only
    consumes persisted final outcomes, so retry counters/errors never leak
    into the DOCX.
    """
    rows: list[list] = []
    for page in pages:
        if not page.source or not page.target:
            continue

        source_data = page.source.get("visual_data") or {}
        target_data = page.target.get("visual_data") or {}
        source_exports = source_data.get("table_exports") or []
        target_exports = target_data.get("table_exports") or []
        source_tables = source_data.get("table_visuals") or []
        target_tables = target_data.get("table_visuals") or []

        order: list[str] = []
        titles: dict[str, str] = {}
        for record in [*source_tables, *target_tables, *source_exports, *target_exports]:
            if not isinstance(record, dict):
                continue
            key = _pairing_key(record)
            if not key:
                continue
            if key not in order:
                order.append(key)
            if not titles.get(key):
                titles[key] = _real_title(record.get("title"))

        def find(records: list, key: str) -> dict | None:
            return next(
                (r for r in records if isinstance(r, dict) and _pairing_key(r) == key),
                None,
            )

        for key in order:
            source_export = find(source_exports, key)
            target_export = find(target_exports, key)
            source_state = _export_state(
                source_export, present=find(source_tables, key) is not None
            )
            target_state = _export_state(
                target_export, present=find(target_tables, key) is not None
            )

            # Only a confirmed success on both sides is PASS.
            if source_state == "yes" and target_state == "yes":
                can_export = "Yes"
                result = _PASS
            elif source_state == "no" or target_state == "no":
                can_export = "No"
                result = _FAIL
            elif source_state == "unknown" or target_state == "unknown":
                # We do not call this PASS or FAIL when the executor did not
                # establish an outcome.
                can_export = ""
                result = ""
            else:
                can_export = ""
                result = ""

            title = titles.get(key)
            if not title:
                title = "Table / Matrix"

            rows.append([
                f"{page.page_name} / {title}",
                can_export,
                _choose(OPT_PASS_FAIL, "Result", result),
            ])
    return rows


# ---------------------------------------------------------------------------
# Report model
# ---------------------------------------------------------------------------


def _nested_value(obj: dict, *paths):
    for path in paths:
        current = obj
        ok = True
        for key in path:
            if not isinstance(current, dict) or key not in current:
                ok = False
                break
            current = current[key]
        if ok and current not in (None, "", [], {}):
            return current
    return ""


def _tester_name(document: dict, response: dict) -> str:
    return _clean(_nested_value(
        document,
        ("details", "tester"),
        ("details", "tester_name"),
        ("execution_context", "tester_name"),
        ("tester_name",),
    )) or _clean(_nested_value(
        response,
        ("details", "tester"),
        ("execution_context", "tester_name"),
        ("tester_name",),
        ("report_metadata", "tester_name"),
    )) or _TESTER_NAME


def _semantic_model_name(document: dict, response: dict) -> str:
    return _clean(_nested_value(
        document,
        ("semantic_model_name",),
        ("details", "semantic_model"),
        ("details", "semantic_model_name"),
        ("report_metadata", "semantic_model_name"),
    )) or _clean(_nested_value(
        response,
        ("semantic_model_name",),
        ("details", "semantic_model"),
        ("details", "semantic_model_name"),
        ("report_metadata", "semantic_model_name"),
    ))


def _validation_date(document: dict, response: dict) -> str:
    return _date_only(
        _nested_value(
            document,
            ("created_at",),
            ("validation_date",),
            ("execution_context", "validation_date"),
        )
        or _nested_value(
            response,
            ("validation_date",),
            ("execution_context", "validation_date"),
            ("report_metadata", "validation_date"),
        )
    )


def _build_model(document: dict, response: dict, run_status: dict) -> dict:
    pages = _collect_pages(document, response)
    ai = response.get("ai_analysis") or {}
    refresh_source = response
    source_executions, target_executions = _split_executions(document)

    report_name = (
        _dom_report_name(response.get("metrics") or [])
        or _dom_report_name(document.get("metrics") or [])
        or _clean(ai.get("dashboard_name"))
        or _clean(response.get("dashboard_name"))
        or _clean(_nested_value(
            response,
            ("report_metadata", "dashboard_name"),
            ("report_metadata", "report_name"),
        ))
    )

    def workspace(executions: list) -> str:
        if not executions:
            return ""
        return _clean((executions[0].get("dashboard") or {}).get("workspace_name"))

    first_source_shot = None
    first_target_shot = None

    def refresh_column(stamp: str, shot: Path | None, source_side: str) -> dict:
        if not stamp:
            return {"stamp": "", "status": None, "screenshot": None, "result": None}

        # If the persisted refresh validation contains an explicit status,
        # use it.  Otherwise a captured refresh stamp is evidence that the
        # refresh field was observed, not proof that a refresh succeeded.
        validation = _nested_value(
            response,
            ("refresh_validation",),
            ("ai_analysis", "refresh_validation"),
        )
        side_data = {}
        comparison_result = ""
        if isinstance(validation, dict):
            side_data = validation.get(source_side) or {}
            comparison = validation.get("comparison") or {}
            comparison_result = _clean(
                comparison.get("result")
                if isinstance(comparison, dict)
                else validation.get("comparison_result")
            )
        elif isinstance(validation, list):
            value_key = "source_value" if source_side == "source" else "target_value"
            for item in validation:
                if not isinstance(item, dict):
                    continue
                if value_key in item:
                    side_data = item
                    comparison_result = _clean(
                        item.get("comparison_result")
                        or item.get("comparison_result_status")
                    )
                    break

        status = _status_label(
            side_data.get("status") if isinstance(side_data, dict) else ""
        )
        result = _status_label(
            side_data.get("result") if isinstance(side_data, dict) else ""
        )
        # Some AI reports store the date comparison result separately.  It is
        # valid evidence for the Result cell, but it is not evidence that the
        # refresh job itself executed successfully.
        if not result and comparison_result:
            result = _status_label(comparison_result)
        if result == _UNCERTAIN:
            result = ""

        screenshot = "Yes" if shot else ""
        return {
            "stamp": stamp,
            "status": status or None,
            "screenshot": screenshot or None,
            "result": result or None,
        }

    source_refresh = _refresh_stamp(source_executions, refresh_source, "source")
    target_refresh = _refresh_stamp(target_executions, refresh_source, "target")
    first_source_shot = _refresh_screenshot(source_executions, "source")
    first_target_shot = _refresh_screenshot(target_executions, "target")

    filter_rows = _filter_rows(pages, document, response)

    visual_rows: list[list] = []
    for page in pages:
        if not page.source and not page.target:
            continue
        visual_rows.extend(_visual_rows(page, response))

    return {
        "report_name": report_name,
        "pages": pages,
        "details": {
            "semantic_model": _semantic_model_name(document, response),
            "source_workspace": workspace(source_executions),
            "target_workspace": workspace(target_executions),
            "validation_date": _validation_date(document, response),
            "tester": _tester_name(document, response),
            "overall": _overall_status(response, run_status),
        },
        "refresh": {
            "source": refresh_column(source_refresh, first_source_shot, "source"),
            "target": refresh_column(target_refresh, first_target_shot, "target"),
            "source_shot": first_source_shot if source_refresh else None,
            "target_shot": first_target_shot if target_refresh else None,
        },
        "filter_rows": filter_rows,
        "visual_rows": visual_rows,
        "export_rows": _export_rows(pages),
    }


# ---------------------------------------------------------------------------
# Word dropdown content controls
# ---------------------------------------------------------------------------

OPT_STATUS = [("Pass✅", "good"), ("Fail❌", "bad"), ("In-Progress🔄", "progress")]
OPT_PASS_FAIL = OPT_STATUS[:2]
OPT_REFRESH = [("Success✅", "good"), ("Failed❌", "bad")]
OPT_YES_NO = [("Yes ✅", "good"), ("No ❌", "bad")]

_CHIP_COLORS = {
    "good": ("1E7B34", RGBColor(0xFF, 0xFF, 0xFF)),
    "bad": ("B00020", RGBColor(0xFF, 0xFF, 0xFF)),
    "progress": ("9DC3E6", RGBColor(0x00, 0x00, 0x00)),
}


@dataclass
class Dropdown:
    options: list[tuple[str, str]]
    selected: int | None
    alias: str


def _choose(options: list, alias: str, key: str | None) -> Dropdown | str:
    """Dropdown pre-selected on the option starting with ``key`` ('' if None)."""
    if key is None:
        return ""
    for index, (label, _kind) in enumerate(options):
        if label.casefold().startswith(key.casefold()):
            return Dropdown(options, index, alias)
    return ""


_sdt_ids = iter(range(1000, 10**7))


def _add_dropdown(paragraph, dropdown: Dropdown, font_pt: float) -> None:
    sdt = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")

    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), dropdown.alias)
    properties.append(alias)
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), dropdown.alias.replace(" ", "_").lower())
    properties.append(tag)
    ident = OxmlElement("w:id")
    ident.set(qn("w:val"), str(next(_sdt_ids)))
    properties.append(ident)
    if dropdown.selected is None:
        properties.append(OxmlElement("w:showingPlcHdr"))

    listing = OxmlElement("w:dropDownList")
    for label, _kind in dropdown.options:
        item = OxmlElement("w:listItem")
        item.set(qn("w:displayText"), label)
        item.set(qn("w:value"), label)
        listing.append(item)
    properties.append(listing)
    sdt.append(properties)

    if dropdown.selected is None:
        text, fill, color_hex, italic = "Choose an item.", None, "808080", True
    else:
        label, kind = dropdown.options[dropdown.selected]
        fill, color = _CHIP_COLORS[kind]
        text, color_hex, italic = label, str(color), False

    content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attribute in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attribute), _FONT)
    run_pr.append(fonts)
    if italic:
        run_pr.append(OxmlElement("w:i"))
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color_hex)
    run_pr.append(color_el)
    for tag_name in ("w:sz", "w:szCs"):
        size = OxmlElement(tag_name)
        size.set(qn("w:val"), str(int(font_pt * 2)))
        run_pr.append(size)
    if fill:
        shade = OxmlElement("w:shd")
        shade.set(qn("w:val"), "clear")
        shade.set(qn("w:color"), "auto")
        shade.set(qn("w:fill"), fill)
        run_pr.append(shade)
    run.append(run_pr)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    run.append(text_el)
    content.append(run)
    sdt.append(content)
    paragraph._p.append(sdt)


# ---------------------------------------------------------------------------
# DOCX layout primitives
# ---------------------------------------------------------------------------


def _set_style_font(style, name: str) -> None:
    style.font.name = name
    fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
    if fonts is not None:  # theme fonts would otherwise override the name
        for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if fonts.get(qn(attribute)) is not None:
                del fonts.attrib[qn(attribute)]


def _configure_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    _set_style_font(normal, _FONT)
    normal.font.size = Pt(_BODY_PT)
    for name, size, before in (("Heading 1", 15, 14), ("Heading 2", 12, 10)):
        style = doc.styles[name]
        _set_style_font(style, _FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = _ACCENT
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.orientation = WD_ORIENT.PORTRAIT
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(_PORTRAIT_MARGIN_IN))


def _shade(cell, fill: str) -> None:
    shade = OxmlElement("w:shd")
    shade.set(qn("w:val"), "clear")
    shade.set(qn("w:color"), "auto")
    shade.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shade)


def _borders(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), _BORDER)
        borders.append(element)
    properties = table._tbl.tblPr
    anchor = None
    for tag in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"):
        anchor = properties.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(borders)
    else:
        properties.append(borders)


def _write(
    cell,
    value,
    *,
    font_pt: float,
    bold: bool = False,
    color: RGBColor | None = None,
    center: bool = False,
) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if isinstance(value, Dropdown):
        _add_dropdown(paragraph, value, font_pt)
        return
    text = "" if value is None else str(value)
    if not text:
        return
    run = paragraph.add_run(text)
    run.font.name = _FONT
    run.font.size = Pt(font_pt)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _add_grid(
    doc: Document,
    headers: list[str] | None,
    rows: list[list],
    widths: list[float],
    *,
    label_column: bool = False,
    font_pt: float = _TABLE_PT,
) -> None:
    table = doc.add_table(rows=0, cols=len(widths))
    table.autofit = False
    _borders(table)

    if headers is not None:
        cells = table.add_row().cells
        for index, text in enumerate(headers):
            _write(cells[index], text, font_pt=font_pt, bold=True, color=_WHITE, center=True)
            _shade(cells[index], _HEADER_FILL)
            cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        row_properties = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        row_properties.append(repeat)

    for values in rows:
        cells = table.add_row().cells
        for index in range(len(widths)):
            value = values[index] if index < len(values) else ""
            is_label = label_column and index == 0
            _write(cells[index], value, font_pt=font_pt, bold=is_label)
            if is_label:
                _shade(cells[index], _LABEL_FILL)
            cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cant_split = OxmlElement("w:cantSplit")
        table.rows[-1]._tr.get_or_add_trPr().append(cant_split)

    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)


def _heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _label(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(10)


def _note(doc: Document, text: str, *, bullet: bool = False) -> None:
    paragraph = doc.add_paragraph(style="List Bullet") if bullet else doc.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = _GRAY


def _add_image(doc: Document, path: Path | None) -> None:
    if not path or not Path(path).is_file():
        return
    try:
        from PIL import Image

        with Image.open(str(path)) as image:
            width_px, height_px = image.size
        scale = min(_IMAGE_MAX_W_IN / width_px, _IMAGE_MAX_H_IN / height_px)
        width, height = Inches(width_px * scale), Inches(height_px * scale)
    except Exception:
        width, height = Inches(_IMAGE_MAX_W_IN), None
    try:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.add_run().add_picture(str(path), width=width, height=height)
    except Exception:
        logger.exception("Failed to embed screenshot | path=%s", path)


def _page_screenshots(doc: Document, prefix: str, pages: list[PagePair]) -> None:
    """GCP then Snowflake shot for each page (the template's order)."""
    for number, page in enumerate(pages, start=1):
        for side, image in (
            (_TARGET_LABEL, page.target_image),
            (_SOURCE_LABEL, page.source_image),
        ):
            if image:
                _label(doc, f"{prefix}.{number} {page.page_name} - {side}")
                _add_image(doc, image)


def _footer_logos(section) -> None:
    logos = [p for p in (_LOGO_LEFT, _LOGO_RIGHT) if p and Path(p).is_file()]
    if not logos:
        return
    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    for cell, path, align in (
        (table.rows[0].cells[0], _LOGO_LEFT, WD_ALIGN_PARAGRAPH.LEFT),
        (table.rows[0].cells[1], _LOGO_RIGHT, WD_ALIGN_PARAGRAPH.RIGHT),
    ):
        paragraph = cell.paragraphs[0]
        paragraph.alignment = align
        if path and Path(path).is_file():
            paragraph.add_run().add_picture(str(path), height=Inches(0.45))


# ---------------------------------------------------------------------------
# Render
# ---------------------------------------------------------------------------


def _render(
    run_id: str,
    document: dict,
    response: dict,
    run_status: dict,
    destination: Path | None,
) -> Path:
    model = _build_model(document, response, run_status)
    pages: list[PagePair] = model["pages"]
    details = model["details"]
    refresh = model["refresh"]
    report_name = model["report_name"]

    doc = Document()
    _configure_document(doc)
    _footer_logos(doc.sections[0])

    # ---- title ------------------------------------------------------------
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(report_name or _DEFAULT_TITLE)
    run.bold = True
    run.underline = True
    run.font.size = Pt(20)
    run.font.color.rgb = _ACCENT
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run(_DOC_SUBTITLE)
    sub_run.font.size = Pt(11)
    sub_run.font.color.rgb = _GRAY

    # ---- 1. Dashboard & Environment Details -------------------------------
    _heading(doc, "1. Dashboard & Environment Details")
    _add_grid(
        doc,
        None,
        [
            ["Dashboard / Report Name", report_name],
            ["Semantic Model Name", details["semantic_model"]],
            [f"{_SOURCE_LABEL} Work Space Name", details["source_workspace"]],
            [f"{_TARGET_LABEL} Work Space Name", details["target_workspace"]],
            ["Tester Name", details["tester"]],
            ["Validation Date", details["validation_date"]],
            ["Overall Status", _choose(OPT_STATUS, "Overall Status", details["overall"])],
        ],
        [2.2, 4.3],
        label_column=True,
        font_pt=_BODY_PT,
    )

    # ---- 2. Semantic Model Refresh Check ----------------------------------
    _heading(doc, "2. Semantic Model Refresh Check")
    source, target = refresh["source"], refresh["target"]

    def refresh_cells(column: dict, side: str) -> list:
        return [
            column.get("stamp", ""),
            _choose(OPT_REFRESH, f"{side} Refresh Status", column.get("status")),
            _choose(OPT_YES_NO, f"{side} Screenshot", column.get("screenshot")),
            _choose(OPT_PASS_FAIL, f"{side} Refresh Result", column.get("result")),
        ]

    labels = ["Refresh Date & Time", "Refresh Status", "Screenshot", "Result"]
    source_cells = refresh_cells(source, _SOURCE_LABEL)
    target_cells = refresh_cells(target, _TARGET_LABEL)
    _add_grid(
        doc,
        ["", _SOURCE_LABEL, _TARGET_REFRESH_HEADER],
        [[labels[i], source_cells[i], target_cells[i]] for i in range(4)],
        [2.0, 2.25, 2.25],
        label_column=True,
        font_pt=_BODY_PT,
    )

    _label(doc, f"2.1 {_TARGET_LABEL}")
    _add_image(doc, refresh["target_shot"])
    _label(doc, f"2.2 {_SOURCE_LABEL}")
    _add_image(doc, refresh["source_shot"])
    _label(doc, "Comments:")

    # ---- 3. Filter Validation Log -----------------------------------------
    doc.add_page_break()
    _heading(doc, "3. Filter Validation Log")
    _note(doc, "Filter/slicer application and landing-page behavior.", bullet=True)
    _add_grid(
        doc,
        ["Page Name", "Action / Filter Applied", "Result"],
        model["filter_rows"] or [["", "", ""]],
        [2.0, 3.0, 1.5],
    )
    # Page screenshots are kept with the visual-validation section below so
    # the same image is not duplicated in the filter section.

    # ---- 4. Visual & Functional Validation Log ----------------------------
    doc.add_page_break()
    _heading(doc, "4. Visual & Functional Validation Log")
    _note(
        doc,
        "Validation of the same visual on both dashboards.",
    )
    _add_grid(
        doc,
        ["Page Name", "Visual Name / Visual Type", "Result"],
        model["visual_rows"] or [["", "", ""]],
        [2.0, 3.5, 1.0],
    )
    _page_screenshots(doc, "4", pages)

    # ---- 5. Data Export Validation Log ------------------------------------
    doc.add_page_break()
    _heading(doc, "5. Data Export Validation Log")
    _note(
        doc,
        "Final table/matrix export outcome. Internal retry attempts are not "
        "reported here.",
    )
    _add_grid(
        doc,
        ["Page / Table", "Can Export?", "Result"],
        model["export_rows"] or [["", "", ""]],
        [3.8, 1.6, 1.1],
    )

    # ---- 6. Summary & Sign-off --------------------------------------------
    _heading(doc, "6. Summary & Sign-off")
    _add_grid(
        doc,
        None,
        [
            ["Sign-Off Status", _choose(OPT_STATUS, "Sign-Off Status", None)],
            ["Tester Name", details["tester"]],
            ["Sign-Off Date", ""],
            ["Signature", ""],
        ],
        [2.2, 4.3],
        label_column=True,
        font_pt=_BODY_PT,
    )

    destination = destination or (
        REPORT_DIR / run_id / report_filename(run_id, report_name)
    )
    destination.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(destination))
    except Exception as exc:
        raise ValidationReportError(
            f"Failed to save DOCX report to {destination}"
        ) from exc

    logger.info(
        "PBI validation report generated | run_id=%s | pages=%d | visual_rows=%d "
        "| export_rows=%d | overall=%s | path=%s",
        run_id, len(pages), len(model["visual_rows"]), len(model["export_rows"]),
        details["overall"], destination,
    )
    return destination


async def build_validation_report(
    run_id: str,
    *,
    destination: Path | None = None,
) -> Path:
    """Generate the validation DOCX from persisted run artifacts.

    No browser is opened here and no new Gemini request is made.  If the
    capture already contains the validation response, that is used directly.
    Otherwise the existing artifact-only validator is used to reconstruct the
    deterministic comparison result.
    """
    document = load_validation_capture(run_id)
    if document is None:
        raise ValidationReportNotFound(
            f"No validation capture found for run_id {run_id}."
        )

    run_status = load_run_status(run_id) or {}

    # Preferred: response persisted alongside the capture.
    response = (
        document.get("validation_response")
        or document.get("validation_report")
        or document.get("comparison_response")
        or document.get("ai_report")
    )

    if not isinstance(response, dict):
        # The existing validator path is retained only as an artifact reader.
        # It must not launch a browser or call Gemini for a report build.
        try:
            from orchestration.validator import DashboardValidator
            response = await DashboardValidator().run_validation_from_artifacts(run_id)
        except Exception as exc:
            logger.exception(
                "Could not reconstruct persisted validation response | run_id=%s",
                run_id,
            )
            raise ValidationReportError(
                f"Persisted validation result is missing for run_id {run_id}. "
                "The report cannot safely invent validation results."
            ) from exc

    if not isinstance(response, dict):
        raise ValidationReportNotFound(
            f"No persisted validation response found for run_id {run_id}."
        )

    return await asyncio.to_thread(
        _render, run_id, document, response, run_status, destination
    )
